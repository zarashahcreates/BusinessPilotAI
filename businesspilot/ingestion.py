"""Secure file ingestion: validation, extraction and OCR.

Uploaded content is treated strictly as untrusted *data*. Nothing here evals, unpickles or
executes anything found in a file. Every loader returns an ``IngestResult`` so the UI can
show an honest note about what was (and was not) extracted.
"""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd

from . import config


@dataclass
class IngestResult:
    filename: str
    table: Optional[pd.DataFrame] = None
    note: str = ""
    level: str = "success"            # success | warning | error
    text_only: bool = False           # True when only free text (no real table) was extracted
    raw_text: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.table is not None and not self.table.empty and not self.text_only


def _ext(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def sanitize_filename(name: str) -> str:
    """Strip any path components and unsafe characters from a user-supplied filename."""
    base = re.split(r"[\\/]", name or "file")[-1]
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", base).strip(" .")
    return base[:120] or "file"


def validate_upload(filename: str, raw: bytes) -> Optional[str]:
    """Return an error message if the upload is unacceptable, else None."""
    ext = _ext(filename)
    if ext not in config.SUPPORTED_EXTENSIONS:
        return f"'.{ext or '?'}' files are not supported. Supported: {', '.join(config.SUPPORTED_EXTENSIONS)}."
    if not raw:
        return "The file is empty."
    if len(raw) > config.MAX_FILE_MB * 1024 * 1024:
        return f"The file is larger than the {config.MAX_FILE_MB} MB limit."
    sigs = config.SIGNATURES.get(ext)
    if sigs and not any(raw.startswith(s) for s in sigs):
        return f"The content does not look like a real .{ext} file (the file may be renamed or corrupted)."
    if ext == "webp" and raw[8:12] != b"WEBP":
        return "The content does not look like a real .webp image."
    return None


def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() if str(c).strip() else f"column_{i+1}" for i, c in enumerate(df.columns)]
    # Make duplicate column names unique instead of silently merging them.
    seen: dict[str, int] = {}
    cols = []
    for c in df.columns:
        if c in seen:
            seen[c] += 1
            cols.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            cols.append(c)
    df.columns = cols
    df = df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)
    return df


def _read_csv(raw: bytes) -> pd.DataFrame:
    last: Exception | None = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(raw), sep=None, engine="python", encoding=enc, nrows=config.MAX_ROWS + 1)
        except UnicodeDecodeError as exc:
            last = exc
        except pd.errors.EmptyDataError:
            raise ValueError("The CSV file has no data.")
    raise ValueError(f"Could not decode the CSV file ({last}).")


def _read_excel(raw: bytes, ext: str) -> pd.DataFrame:
    engine = "openpyxl" if ext == "xlsx" else "xlrd"
    try:
        sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None, engine=engine)
    except ImportError:
        if ext == "xls":
            raise RuntimeError("Reading legacy .xls files needs the 'xlrd' package (pip install xlrd) "
                               "— or re-save the file as .xlsx.")
        raise
    frames = []
    for name, sheet in sheets.items():
        sheet = _clean_frame(sheet)
        if not sheet.empty:
            if len(sheets) > 1:
                sheet["sheet"] = name
            frames.append(sheet)
    if not frames:
        raise ValueError("The workbook contains no data.")
    return pd.concat(frames, ignore_index=True, sort=False)


def _read_json(raw: bytes) -> pd.DataFrame:
    try:
        obj = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid JSON: {exc}")
    if isinstance(obj, dict):
        # Prefer the first list-of-records value, otherwise treat the object as one record.
        for v in obj.values():
            if isinstance(v, list) and v and isinstance(v[0], dict):
                obj = v
                break
        else:
            obj = [obj]
    if not isinstance(obj, list) or not obj:
        raise ValueError("JSON must contain a list of records or an object with a list of records.")
    return pd.json_normalize(obj)


def _text_to_table(text: str) -> Optional[pd.DataFrame]:
    """Try to interpret plain text as delimited data. Returns None if it isn't clearly tabular."""
    lines = [l for l in text.splitlines() if l.strip()]
    if len(lines) < 2:
        return None
    for sep in (",", ";", "\t", "|"):
        counts = [l.count(sep) for l in lines[:20]]
        if counts[0] >= 1 and len(set(counts)) == 1:
            try:
                df = pd.read_csv(io.StringIO("\n".join(lines)), sep=sep, engine="python")
                if df.shape[1] >= 2:
                    return df
            except Exception:
                return None
    return None


def _read_txt(raw: bytes) -> tuple[str, Optional[pd.DataFrame]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    return text, _text_to_table(text)


def _read_pdf(raw: bytes) -> tuple[Optional[pd.DataFrame], str]:
    import pdfplumber
    frames, texts = [], []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        pages = pdf.pages[: config.MAX_PDF_PAGES]
        for page in pages:
            for table in page.extract_tables() or []:
                if table and len(table) > 1:
                    header = [str(x or f"column_{i+1}") for i, x in enumerate(table[0])]
                    frames.append(_clean_frame(pd.DataFrame(table[1:], columns=header)))
            texts.append(page.extract_text() or "")
    frame = None
    if frames:
        frame = pd.concat(frames, ignore_index=True, sort=False)
    return frame, "\n".join(texts)


def _read_docx(raw: bytes) -> tuple[Optional[pd.DataFrame], str]:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    frames = []
    for t in doc.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if len(rows) > 1:
            frames.append(_clean_frame(pd.DataFrame(rows[1:], columns=rows[0])))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    frame = pd.concat(frames, ignore_index=True, sort=False) if frames else None
    return frame, text


def ocr_available() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image(raw: bytes) -> str:
    from PIL import Image
    import pytesseract
    img = Image.open(io.BytesIO(raw))
    img.load()
    if img.width * img.height > 40_000_000:
        raise ValueError("The image is too large to process safely.")
    return pytesseract.image_to_string(img)


def _numeric_text_to_table(text: str) -> Optional[pd.DataFrame]:
    """Parse OCR / PDF text lines such as 'Widget   12   450.00' into a table, if consistent."""
    rows = []
    for line in text.splitlines():
        parts = re.split(r"\s{2,}|\t|\s\|\s", line.strip())
        if len(parts) >= 2:
            rows.append(parts)
    if len(rows) < 3:
        return None
    width = max(set(len(r) for r in rows), key=[len(r) for r in rows].count)
    rows = [r for r in rows if len(r) == width]
    if len(rows) < 3 or width < 2:
        return None
    return pd.DataFrame(rows[1:], columns=[str(c) or f"column_{i+1}" for i, c in enumerate(rows[0])])


def load_file(filename: str, raw: bytes) -> IngestResult:
    """Validate and extract a single uploaded file. Never raises; failures are reported in the result."""
    filename = sanitize_filename(filename)
    err = validate_upload(filename, raw)
    if err:
        return IngestResult(filename, note=err, level="error")
    ext = _ext(filename)
    try:
        if ext == "csv":
            df = _clean_frame(_read_csv(raw))
        elif ext in ("xlsx", "xls"):
            df = _read_excel(raw, ext)
        elif ext == "json":
            df = _clean_frame(_read_json(raw))
        elif ext == "txt":
            text, df = _read_txt(raw)
            if df is None:
                return IngestResult(filename, note="Plain text was read but does not look like a table "
                                    "(use CSV-style rows with a header line, or enter data manually).",
                                    level="warning", text_only=True, raw_text=text)
            df = _clean_frame(df)
        elif ext == "pdf":
            df, text = _read_pdf(raw)
            if df is None:
                df = _numeric_text_to_table(text) if text.strip() else None
                if df is None:
                    if not text.strip():
                        msg = ("No readable text found. This looks like a scanned PDF. "
                               "Convert pages to images (PNG/JPG) and upload them for OCR"
                               + ("." if ocr_available() else ", and install Tesseract OCR first (see README)."))
                        return IngestResult(filename, note=msg, level="warning", text_only=True)
                    return IngestResult(filename, note="Text was found but no reliable table. "
                                        "Upload the source spreadsheet if you have it.", level="warning",
                                        text_only=True, raw_text=text)
                return IngestResult(filename, table=_clean_frame(df),
                                    note="Table reconstructed from PDF text layout — please verify the values.",
                                    level="warning", raw_text=text)
            df = _clean_frame(df)
        elif ext == "docx":
            df, text = _read_docx(raw)
            if df is None:
                df = _text_to_table(text)
                if df is None:
                    return IngestResult(filename, note="The document has no tables to analyse.",
                                        level="warning", text_only=True, raw_text=text)
            df = _clean_frame(df)
        elif ext == "doc":
            return IngestResult(filename, level="error",
                                note="Legacy .doc files cannot be read reliably. Please save the file as .docx and upload it again.")
        else:  # images
            if not ocr_available():
                return IngestResult(filename, level="warning", text_only=True,
                                    note="OCR is not available: install the Tesseract OCR program (see README) to read images.")
            text = _ocr_image(raw)
            df = _numeric_text_to_table(text)
            if df is None:
                return IngestResult(filename, level="warning", text_only=True, raw_text=text,
                                    note="Text was detected but could not be turned into a reliable table. "
                                         "Use a clearer image or enter the figures manually.")
            return IngestResult(filename, table=_clean_frame(df), level="warning", raw_text=text,
                                note="Table reconstructed with OCR — OCR can misread digits, please verify.")
    except Exception as exc:  # noqa: BLE001 - convert any parser failure to a friendly message
        return IngestResult(filename, level="error", note=f"Could not read this file: {type(exc).__name__}: {str(exc)[:200]}")

    if df.empty:
        return IngestResult(filename, level="error", note="No data rows were found in this file.")
    warnings = []
    if len(df) > config.MAX_ROWS:
        df = df.head(config.MAX_ROWS)
        warnings.append(f"Only the first {config.MAX_ROWS:,} rows were loaded.")
    return IngestResult(filename, table=df, note=f"Loaded {len(df):,} rows and {df.shape[1]} columns.", warnings=warnings)


def parse_manual_data(text: str) -> pd.DataFrame:
    if not text or not text.strip():
        raise ValueError("Please enter some CSV data.")
    if len(text) > 2_000_000:
        raise ValueError("Manual data is too large; upload a file instead.")
    df = _clean_frame(pd.read_csv(io.StringIO(text.strip()), sep=None, engine="python"))
    if df.empty:
        raise ValueError("The supplied data has no rows.")
    return df


def combine_tables(results: list[IngestResult]) -> pd.DataFrame:
    """Stack tables from several files; columns with the same name line up, a source column is added."""
    frames = []
    for r in results:
        if r.ok:
            t = r.table.copy()
            t["source_file"] = r.filename
            frames.append(t)
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True, sort=False)
